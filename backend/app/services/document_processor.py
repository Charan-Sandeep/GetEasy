"""
Extracts raw text from uploaded academic documents (PDF, DOCX, PPTX)
and splits it into chunks ready for embedding.
"""
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
import re


def extract_text(file_path: str, filename: str) -> str:
    ext = filename.lower().split(".")[-1]

    if ext == "pdf":
        return _extract_pdf(file_path)
    elif ext == "docx":
        return _extract_docx(file_path)
    elif ext == "pptx":
        return _extract_pptx(file_path)
    elif ext == "txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _extract_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _extract_pptx(file_path: str) -> str:
    prs = Presentation(file_path)
    slides_text = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        slides_text.append(f"[Slide {i}]\n" + "\n".join(texts))
    return "\n\n".join(slides_text)


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 150) -> list[str]:
    """
    Simple sliding-window chunker by character count, snapped to sentence
    boundaries where possible. Good enough for Week 1-2 — swap for a
    smarter semantic chunker later if needed.
    """
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        # try to snap to the nearest sentence end
        if end < len(text):
            snap = text.rfind(". ", start, end)
            if snap != -1 and snap > start:
                end = snap + 1
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start = end - overlap if end - overlap > start else end

    return chunks


def guess_content_type(filename: str, text_sample: str) -> str:
    """
    Cheap heuristic tagger for Week 1-2. Replace with a real classifier
    once the pipeline works end-to-end.
    """
    name = filename.lower()
    sample = text_sample.lower()

    if any(k in name for k in ["question", "qp", "paper", "exam"]):
        return "questions"
    if any(k in name for k in ["lab", "manual", "practical"]):
        return "labs"
    if any(k in name for k in ["note", "lecture", "slide"]):
        return "notes"
    if any(k in sample[:500] for k in ["question 1", "q1.", "attempt any"]):
        return "questions"
    return "textbook"


def extract_numbered_questions(text: str, limit: int = 20) -> list[str]:
    """Split a question paper into numbered questions for batch solving."""
    marker = re.compile(r"(?im)(?=^\s*(?:q(?:uestion)?\s*)?\d{1,3}\s*[.)])")
    starts = list(marker.finditer(text))
    if not starts:
        return []
    questions = []
    for index, match in enumerate(starts):
        end = starts[index + 1].start() if index + 1 < len(starts) else len(text)
        question = re.sub(r"\s+", " ", text[match.start():end]).strip()
        if len(question) >= 12:
            questions.append(question[:1800])
        if len(questions) >= limit:
            break
    return questions
